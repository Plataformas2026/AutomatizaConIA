"""
Lectura LOCAL de Excel y CSV, en memoria. El contenido del archivo vive solo dentro de estas
funciones y de engine.py: nunca se escribe en disco, ni en logs, ni se envía a ningún sitio.
Replica src/lib/parsing/parse.ts de la web.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from typing import Any, Iterable, Iterator, Optional

from engine import Row, SheetTable, Table, resolve_sheets

MAX_ROWS = 50_000
MAX_SHEETS = 50
MAX_COLUMNS = 500
EXCEL_EXTS = (".xlsx", ".xlsm", ".xls")
SUPPORTED_EXTS = EXCEL_EXTS + (".csv",)


class ReadError(Exception):
    """El archivo no se puede leer (dañado, protegido con contraseña, formato raro…).
    El mensaje NUNCA incluye contenido del archivo."""


@dataclass
class Workbook:
    available: list[str]                      # todas las hojas ([] en CSV)
    structure: dict                           # {"sheets": [...], "tables": [{"sheet":…, "columns":[…]}]} (solo nombres)
    sheets: list[SheetTable] = field(default_factory=list)  # solo las hojas que toca evaluar


def _is_blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def _normalize_headers(raw: list[Any]) -> list[str]:
    last = len(raw) - 1
    while last >= 0 and _is_blank(raw[last]):
        last -= 1
    seen: dict[str, int] = {}
    out: list[str] = []
    for i, h in enumerate(raw[: last + 1]):
        base = f"Columna {i + 1}" if _is_blank(h) else str(h).strip()
        seen[base] = seen.get(base, 0) + 1
        out.append(base if seen[base] == 1 else f"{base} ({seen[base]})")
    return out[:MAX_COLUMNS]


def _clean(v: Any) -> Any:
    """Solo tipos que entiende el motor: str, número, bool, fecha o None."""
    if v is None or v == "":
        return None
    if isinstance(v, (str, bool, int, float, datetime, date)):
        return v
    if isinstance(v, (time, timedelta)):
        return str(v)
    return str(v)


def _first_used(arr: list[Any]) -> int:
    for i, c in enumerate(arr):
        if not _is_blank(c):
            return i
    return len(arr)


def _build_table(rows: Iterable[tuple[int, list[Any]]], header_only: bool = False) -> Table:
    """La 1ª fila con datos es la cabecera; las filas vacías se saltan.
    Las columnas vacías a la izquierda se ignoran (como hace la web: el rango empieza en la 1ª columna usada)."""
    raw: list[tuple[int, list[Any]]] = []
    for number, arr in rows:
        if all(_is_blank(c) for c in arr):
            continue
        raw.append((number, arr))
        if header_only or len(raw) > MAX_ROWS:
            break
    if not raw:
        return Table([], [])

    lead = min(_first_used(arr) for _, arr in raw)
    header_arr = raw[0][1][lead:]
    headers = _normalize_headers(header_arr)
    if header_only:
        return Table(headers, [])

    out: list[Row] = []
    for number, arr in raw[1:]:
        cells = arr[lead:]
        out.append(Row(number, {h: _clean(cells[j]) if j < len(cells) else None for j, h in enumerate(headers)}))
    return Table(headers, out)


# ── CSV ─────────────────────────────────────────────────────────────────
def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")  # CSV de Excel en Windows


def _detect_delimiter(text: str) -> str:
    lines = text.splitlines()[:20]
    best, best_score = ",", (0, 0)
    for d in (",", ";", "\t", "|"):
        counts = [len(r) for r in csv.reader(lines, delimiter=d) if r]
        if not counts or counts[0] < 2:
            continue
        score = (sum(1 for c in counts if c == counts[0]), counts[0])
        if score > best_score:
            best, best_score = d, score
    return best


def _csv_rows(text: str) -> Iterator[tuple[int, list[Any]]]:
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=_detect_delimiter(text))
    for i, row in enumerate(reader, start=1):
        yield i, row


def _read_csv(data: bytes) -> Workbook:
    table = _build_table(_csv_rows(_decode(data)))
    return Workbook([], {"sheets": [], "tables": [{"sheet": None, "columns": table.headers}]}, [SheetTable(None, table)])


# ── XLSX / XLSM ─────────────────────────────────────────────────────────
def _read_xlsx(data: bytes, mode: str, names: list[str]) -> Workbook:
    from openpyxl import load_workbook  # import diferido: arranque más rápido

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # zip dañado, contraseña, etc.
        raise ReadError("No se pudo abrir el Excel (¿está dañado o protegido con contraseña?)") from type(e)

    try:
        available = list(wb.sheetnames)
        wanted = set(resolve_sheets(available, mode, names))
        tables: list[dict] = []
        selected: list[SheetTable] = []
        for name in available[:MAX_SHEETS]:
            ws = wb[name]
            if not hasattr(ws, "iter_rows"):  # hoja de gráfico
                continue
            rows = ((i, list(r)) for i, r in enumerate(ws.iter_rows(values_only=True), start=1))
            table = _build_table(rows, header_only=name not in wanted)
            tables.append({"sheet": name, "columns": table.headers})
            if name in wanted:
                selected.append(SheetTable(name, table))
        return Workbook(available, {"sheets": available, "tables": tables}, selected)
    finally:
        wb.close()


# ── XLS (Excel antiguo) ─────────────────────────────────────────────────
def _read_xls(data: bytes, mode: str, names: list[str]) -> Workbook:
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
    except Exception as e:
        raise ReadError("No se pudo abrir el Excel (¿está dañado o protegido con contraseña?)") from type(e)

    try:
        available = book.sheet_names()
        wanted = set(resolve_sheets(available, mode, names))
        tables: list[dict] = []
        selected: list[SheetTable] = []
        for name in available[:MAX_SHEETS]:
            sh = book.sheet_by_name(name)

            def rows(sh=sh) -> Iterator[tuple[int, list[Any]]]:
                for r in range(sh.nrows):
                    vals: list[Any] = []
                    for cell in sh.row(r):
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            try:
                                vals.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                            except Exception:
                                vals.append(None)
                        elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                            vals.append(bool(cell.value))
                        elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                            vals.append(None)
                        else:
                            vals.append(cell.value)
                    yield r + 1, vals

            table = _build_table(rows(), header_only=name not in wanted)
            tables.append({"sheet": name, "columns": table.headers})
            if name in wanted:
                selected.append(SheetTable(name, table))
        return Workbook(available, {"sheets": available, "tables": tables}, selected)
    finally:
        book.release_resources()


# ── API pública ─────────────────────────────────────────────────────────
def read_workbook(data: bytes, extension: str, sheet_mode: str = "first", sheet_names: Optional[list[str]] = None) -> Workbook:
    ext = extension.lower()
    names = sheet_names or []
    if ext == ".csv":
        return _read_csv(data)
    if ext in (".xlsx", ".xlsm"):
        return _read_xlsx(data, sheet_mode, names)
    if ext == ".xls":
        return _read_xls(data, sheet_mode, names)
    raise ReadError("Tipo de archivo no admitido")
